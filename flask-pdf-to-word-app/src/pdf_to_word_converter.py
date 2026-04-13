import io
import os

import pytesseract
from docx import Document
from docx.shared import Inches
from pdf2docx import Converter
from pdf2image import convert_from_path

try:
    import fitz  # PyMuPDF
    _HAS_PYMUPDF = True
except ImportError:
    _HAS_PYMUPDF = False


def _extract_page_images(pdf_path: str, page_index: int) -> list:
    """Extract embedded raster images from a PDF page (0-indexed). Returns list of BytesIO."""
    if not _HAS_PYMUPDF:
        return []
    images = []
    try:
        doc = fitz.open(pdf_path)
        page = doc[page_index]
        for img_info in page.get_images(full=True):
            xref = img_info[0]
            base = doc.extract_image(xref)
            images.append(io.BytesIO(base['image']))
        doc.close()
    except Exception:
        pass
    return images


def _ocr_with_layout(image, lang: str) -> list:
    """
    Use pytesseract bounding-box data to reconstruct paragraphs with heading detection.
    Returns a list of (text, is_heading) tuples.
    """
    try:
        data = pytesseract.image_to_data(
            image, lang=lang, output_type=pytesseract.Output.DICT
        )
    except Exception:
        return [(pytesseract.image_to_string(image, lang=lang), False)]

    # Derive a median word height as a baseline for heading detection
    word_heights = [
        data['height'][i]
        for i in range(len(data['text']))
        if data['level'][i] == 5
        and data['conf'][i] != -1
        and data['text'][i].strip()
    ]
    if not word_heights:
        return [('', False)]
    sorted_h = sorted(word_heights)
    median_h = sorted_h[len(sorted_h) // 2]

    # Aggregate words → lines
    lines: dict = {}
    for i in range(len(data['text'])):
        if data['level'][i] != 5:
            continue
        word = data['text'][i].strip()
        if not word or data['conf'][i] == -1:
            continue
        key = (data['block_num'][i], data['par_num'][i], data['line_num'][i])
        if key not in lines:
            lines[key] = {'words': [], 'max_h': 0}
        lines[key]['words'].append(word)
        lines[key]['max_h'] = max(lines[key]['max_h'], data['height'][i])

    # Aggregate lines → paragraphs
    paras: dict = {}
    for (blk, par, ln), val in sorted(lines.items()):
        key = (blk, par)
        if key not in paras:
            paras[key] = {'lines': [], 'max_h': 0}
        paras[key]['lines'].append(' '.join(val['words']))
        paras[key]['max_h'] = max(paras[key]['max_h'], val['max_h'])

    result = []
    for key in sorted(paras):
        val = paras[key]
        text = ' '.join(val['lines'])
        is_heading = median_h > 0 and val['max_h'] > median_h * 1.5
        result.append((text, is_heading))

    return result or [('', False)]


def pdf_to_word(pdf_path, word_path, lang='eng', file_id=None, progress=None,
                page_from=None, page_to=None):
    def update(msg: str) -> None:
        if progress is not None and file_id is not None:
            progress[file_id] = msg

    update('Opening document...')

    # --- Attempt direct conversion via pdf2docx ---
    try:
        cv = Converter(pdf_path)
        update('Analyzing document...')
        cv_kwargs = {}
        if page_from is not None:
            cv_kwargs['start'] = page_from - 1  # pdf2docx: 0-indexed, inclusive start
        if page_to is not None:
            cv_kwargs['end'] = page_to           # pdf2docx: 0-indexed, exclusive end
        cv.convert(word_path, **cv_kwargs)
        cv.close()
        if os.path.exists(word_path) and os.path.getsize(word_path) > 0:
            update('Conversion complete!')
            return
    except Exception:
        update('Direct conversion failed, switching to OCR...')

    # --- OCR fallback ---
    update('Converting pages to images...')
    img_kwargs = {}
    if page_from is not None:
        img_kwargs['first_page'] = page_from  # pdf2image: 1-indexed, inclusive
    if page_to is not None:
        img_kwargs['last_page'] = page_to
    images = convert_from_path(pdf_path, **img_kwargs)

    doc = Document()
    first_page_0idx = (page_from - 1) if page_from is not None else 0

    for i, image in enumerate(images):
        update(f'OCR page {i + 1}/{len(images)}')

        # Embed raster images extracted from this PDF page
        for img_bytes in _extract_page_images(pdf_path, first_page_0idx + i):
            try:
                doc.add_picture(img_bytes, width=Inches(5))
            except Exception:
                pass

        # OCR with layout-aware paragraph and heading detection
        for text, is_heading in _ocr_with_layout(image, lang):
            if not text.strip():
                continue
            if is_heading:
                doc.add_heading(text, level=2)
            else:
                doc.add_paragraph(text)

        if i < len(images) - 1:
            doc.add_page_break()

    doc.save(word_path)
    update('OCR complete!')